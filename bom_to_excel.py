#!/usr/bin/env python3
"""Turn a BOM CSV into a formatted Excel quote sheet with a summary tab."""
import argparse, csv, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REQUIRED = ["part_number", "description", "qty", "unit_cost"]
HDR = PatternFill("solid", fgColor="1F4E78"); WHITE = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF"); BOX = Border(THIN, THIN, THIN, THIN)

def read_bom(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        r = csv.DictReader(f)
        cols = [c.strip().lower() for c in (r.fieldnames or [])]
        missing = [c for c in REQUIRED if c not in cols]
        if missing: sys.exit(f"Missing columns: {', '.join(missing)}")
        rows, errors = [], []
        for i, raw in enumerate(r, start=2):
            row = {k.strip().lower(): (v or "").strip() for k, v in raw.items() if k}
            try:
                row["qty"] = float(row["qty"]); row["unit_cost"] = float(row["unit_cost"])
            except ValueError:
                errors.append(f"line {i}: qty/unit_cost not a number, skipped"); continue
            rows.append(row)
    return rows, errors

def header(ws, cols, row=1):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.fill, cell.font, cell.border = HDR, WHITE, BOX
        cell.alignment = Alignment(horizontal="center")

def build(rows, out, project, currency, markup, tax):
    wb = Workbook(); ws = wb.active; ws.title = "Quote"
    cols = ["#", "Part No.", "Description", "Material", "Process", "Qty",
            f"Unit cost ({currency})", f"Line total ({currency})"]
    ws["A1"] = f"Quote: {project}"; ws["A1"].font = Font(bold=True, size=14)
    header(ws, cols, 3)
    money = '#,##0.00'
    for i, r in enumerate(rows, 1):
        n = i + 3
        vals = [i, r["part_number"], r["description"], r.get("material", ""),
                r.get("process", ""), r["qty"], r["unit_cost"], f"=F{n}*G{n}"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=n, column=j, value=v); c.border = BOX
            if j in (7, 8): c.number_format = money
            if i % 2 == 0: c.fill = PatternFill("solid", fgColor="EAF1FB")
    last = len(rows) + 3; t = last + 2
    lines = [("Subtotal", f"=SUM(H4:H{last})"),
             (f"Markup ({markup:.0%})", f"=H{t}*{markup}"),
             (f"Tax ({tax:.0%})", f"=(H{t}+H{t+1})*{tax}"),
             ("Grand total", f"=H{t}+H{t+1}+H{t+2}")]
    for k, (lab, f) in enumerate(lines):
        ws.cell(row=t+k, column=7, value=lab).font = Font(bold=True)
        c = ws.cell(row=t+k, column=8, value=f); c.number_format = money; c.font = Font(bold=k == 3)
    for j, w in enumerate([5, 12, 34, 16, 18, 7, 16, 18], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"

    s = wb.create_sheet("Summary"); s["A1"] = f"Cost summary: {project}"; s["A1"].font = Font(bold=True, size=14)
    by = defaultdict(lambda: [0, 0.0])
    for r in rows:
        k = r.get("process") or "Unspecified"; by[k][0] += r["qty"]; by[k][1] += r["qty"] * r["unit_cost"]
    header(s, ["Process", "Total qty", f"Cost ({currency})", "Share"], 3)
    total = sum(v[1] for v in by.values()) or 1
    for i, (k, (q, c)) in enumerate(sorted(by.items(), key=lambda x: -x[1][1]), 4):
        for j, v in enumerate([k, q, c, c / total], 1):
            cell = s.cell(row=i, column=j, value=v); cell.border = BOX
        s.cell(row=i, column=3).number_format = money; s.cell(row=i, column=4).number_format = "0.0%"
    n = len(by) + 5
    s.cell(row=n, column=1, value="Unique parts").font = Font(bold=True); s.cell(row=n, column=2, value=len(rows))
    s.cell(row=n+1, column=1, value="Total pieces").font = Font(bold=True); s.cell(row=n+1, column=2, value=sum(r["qty"] for r in rows))
    s.cell(row=n+2, column=1, value="Grand total").font = Font(bold=True)
    c = s.cell(row=n+2, column=2, value=f"=Quote!H{t+3}"); c.number_format = money
    for j, w in enumerate([22, 12, 16, 10], 1): s.column_dimensions[get_column_letter(j)].width = w
    for sh in (ws, s):
        sh.page_setup.orientation = "landscape"; sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 1
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("csv"); p.add_argument("-o", "--output", default="quote.xlsx")
    p.add_argument("--project", default="BOM quote"); p.add_argument("--currency", default="INR")
    p.add_argument("--markup", type=float, default=0.15); p.add_argument("--tax", type=float, default=0.18)
    a = p.parse_args()
    rows, errors = read_bom(a.csv)
    for e in errors: print("Warning:", e)
    if not rows: sys.exit("No valid rows found.")
    build(rows, a.output, a.project, a.currency, a.markup, a.tax)
    print(f"Wrote {a.output}: {len(rows)} parts")

if __name__ == "__main__":
    main()
